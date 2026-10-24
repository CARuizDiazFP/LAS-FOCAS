# Nombre de archivo: linea_base_busqueda_camaras.py
# Ubicación de archivo: scripts/linea_base_busqueda_camaras.py
# Descripción: Graba y compara la línea base de smart search y exportación de Cámaras contra la base de dev
"""Línea base de la smart search y la exportación de Cámaras.

Sirve para verificar que un refactor de la smart search (`POST /api/infra/smart-search`)
o de la exportación (`GET /api/infra/export/cameras`) de Cámaras no cambia las respuestas.

Sólo lecturas (SELECT). Variables de entorno para la base de DEV:

    POSTGRES_HOST=127.0.0.1 POSTGRES_PORT=5433 POSTGRES_DB=focas_dev \
    POSTGRES_USER=FOCALBOT POSTGRES_PASSWORD="$(cat .secrets/Dev_db_password_v1.txt)"

Aborta con exit 2 si no es host local, puerto 5433 y base focas_dev (nunca producción, 5432).

Uso:

    python scripts/linea_base_busqueda_camaras.py grabar --salida DIR [--solo smart|export]
    python scripts/linea_base_busqueda_camaras.py comparar --salida DIR [--solo smart|export]
    python scripts/linea_base_busqueda_camaras.py medir --salida DIR [--solo smart|export]

`medir` no compara respuestas: por cada consulta registra el tiempo de pared y la cantidad de
sentencias SQL (listener `before_cursor_execute` sobre `db.session.engine`, el que usa
`SessionLocal`), con las sentencias más repetidas. Escribe `DIR/tiempos.json` e imprime el
total, la mediana y el peor caso por grupo.

Tiempos aproximados: smart unos 20 min, export unos 1,5 min.

Advertencia de deriva: la base de dev cambia (los workers escriben), así que la línea base
se graba FRESCA con el código anterior justo antes de comparar. Pasos:

    1. git worktree add --detach <scratch>/base <commit-anterior>
    2. Enlazar `.venv` y `.secrets` dentro de ese worktree.
    3. `grabar` en <scratch>/base.
    4. `comparar` en el worktree con el cambio.
    5. Si difiere, repetir una vez para descartar deriva.
    6. Al final, `git worktree remove <scratch>/base` sin `--force`.
"""

from __future__ import annotations

import argparse
import contextlib
import csv
import difflib
import io
import json
import os
import re
import statistics
import sys
import time
from collections import Counter
from pathlib import Path
from unittest.mock import patch

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

PUERTO_DEV = "5433"
BASE_DEV = "focas_dev"
HOSTS_LOCALES = {"127.0.0.1", "localhost", "::1"}
_TS = re.compile(r"camaras_\d{8}_\d{6}")
TOP_SENTENCIAS = 5

# Lista de mediciones de `medir`; None cuando se graba o compara (el medidor no hace nada).
_MEDICIONES: list[dict] | None = None


@contextlib.contextmanager
def _medicion(nombre: str):
    """Mide tiempo y sentencias SQL de un bloque; cede el registro para agregarle datos DESPUÉS del bloque
    (fuera del tiempo medido)."""
    if _MEDICIONES is None:
        yield {}
        return
    from sqlalchemy import event

    from db.session import engine

    conteo: Counter[str] = Counter()
    ejemplos: dict[str, str] = {}

    def _contar(conn, cursor, statement, parameters, context, executemany):
        conteo[statement] += 1
        ejemplos.setdefault(statement, repr(parameters))

    registro: dict = {"nombre": nombre}
    event.listen(engine, "before_cursor_execute", _contar)
    inicio = time.perf_counter()
    try:
        yield registro
    finally:
        segundos = time.perf_counter() - inicio
        event.remove(engine, "before_cursor_execute", _contar)
        registro["segundos"] = round(segundos, 3)
        registro["sentencias"] = sum(conteo.values())
        registro["mas_repetidas"] = [
            {"veces": n, "sql": sql, "parametros_ejemplo": ejemplos[sql]} for sql, n in conteo.most_common(TOP_SENTENCIAS)
        ]
        _MEDICIONES.append(registro)


def _abortar(mensaje: str) -> None:
    print(f"ABORTADO: {mensaje}", file=sys.stderr)
    sys.exit(2)


def _verificar_base_dev() -> None:
    """Verifica variables de entorno y la URL efectiva del engine (sin imprimir la contraseña)."""
    host = os.getenv("POSTGRES_HOST", "")
    puerto = os.getenv("POSTGRES_PORT", "")
    if host not in HOSTS_LOCALES or puerto != PUERTO_DEV:
        _abortar(f"se exige POSTGRES_HOST local y POSTGRES_PORT={PUERTO_DEV} (recibido host={host!r} puerto={puerto!r})")
    from db.session import engine

    url = engine.url
    if str(url.port) != PUERTO_DEV or url.host not in HOSTS_LOCALES or url.database != BASE_DEV:
        _abortar(
            f"la URL efectiva apunta a {url.host}:{url.port}/{url.database}; "
            f"sólo se permite {BASE_DEV} local en {PUERTO_DEV}"
        )


def _combinaciones_smart() -> list[tuple[str, dict]]:
    terminos = {
        "vacio": [],
        "ruta": ["ruta"],
        "ruta_5_espacios": ["RUTA ", " 5"],
        "repetido": ["a", "a"],
        "cadena_vacia": [""],
    }
    estados = {"sin": None, "libre": "LIBRE", "nooper": "NO_OPERATIVA"}
    consultas = []
    for tn, t in terminos.items():
        for en, e in estados.items():
            for inc in (False, True):
                for off in (0, 50):
                    nombre = f"smart_{tn}_est-{en}_inc-{int(inc)}_off-{off}"
                    cuerpo = {"terms": t, "estado": e, "incluir_no_operativas": inc, "limit": 50, "offset": off}
                    consultas.append((nombre, cuerpo))
    return consultas


def _servicio_con_camaras() -> str:
    """Primer servicio (menor id) con alguna ruta activa cuyo empalme tenga cámara (sólo SELECT)."""
    from sqlalchemy import text

    from db.session import SessionLocal

    consulta = text(
        "SELECT s.servicio_id FROM app.servicios s "
        "JOIN app.rutas_servicio r ON r.servicio_id = s.id AND r.activa "
        "JOIN app.ruta_empalme_association a ON a.ruta_id = r.id "
        "JOIN app.empalmes e ON e.id = a.empalme_id AND e.camara_id IS NOT NULL "
        "WHERE s.servicio_id IS NOT NULL ORDER BY s.id LIMIT 1"
    )
    with SessionLocal() as sesion:
        fila = sesion.execute(consulta).first()
    if fila is None:
        _abortar("no se encontró ningún servicio con cámaras en la base de dev")
    return fila[0]


def _normalizar_ts(texto: str) -> str:
    return _TS.sub("camaras_TS", texto)


def _xlsx_crudo(contenido: bytes) -> str:
    """Lee el XLSX con openpyxl (celdas crudas, sin inferencia de tipos): encabezado + filas."""
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hojas = {}
    for hoja in libro.worksheets:
        filas = [list(fila) for fila in hoja.iter_rows(values_only=True)]
        hojas[hoja.title] = {"columnas": filas[0] if filas else [], "filas": filas[1:]}
    libro.close()
    return json.dumps(hojas, ensure_ascii=False, indent=1, default=repr)


def _correr_smart(cliente, resultados: dict[str, str]) -> None:
    for nombre, cuerpo in _combinaciones_smart():
        with _medicion(nombre) as registro:
            r = cliente.post("/api/infra/smart-search", json=cuerpo)
        if _MEDICIONES is not None:
            registro["camaras"] = len(r.json().get("camaras", [])) if r.status_code == 200 else 0
        resultados[nombre] = json.dumps({"status": r.status_code, "cuerpo": r.json()}, ensure_ascii=False, indent=1)


def _filas_exportadas(respuesta, formato: str, disposicion: str) -> int:
    """Cantidad de cámaras (filas de datos) de una exportación; 0 si no fue 200. Sólo se usa en `medir`."""
    if respuesta.status_code != 200:
        return 0
    if formato == "xlsx" and "xlsx" in disposicion:
        hojas = json.loads(_xlsx_crudo(respuesta.content))
        return sum(len(h["filas"]) for h in hojas.values())
    return max(len(list(csv.reader(io.StringIO(respuesta.content.decode("utf-8"))))) - 1, 0)


def _correr_export(cliente, resultados: dict[str, str]) -> None:
    servicio_id = _servicio_con_camaras()
    resultados["_servicio_id"] = servicio_id
    for estado in ("ALL", "LIBRE", "BANEADA"):
        for formato in ("csv", "xlsx"):
            base = {"filter_status": estado, "format": formato}
            for variante, p in (("", base), ("_srv", {**base, "servicio_id": servicio_id})):
                with _medicion(f"export_{estado}_{formato}{variante}") as registro:
                    r = cliente.get("/api/infra/export/cameras", params=p)
                disp = _normalizar_ts(r.headers.get("content-disposition", ""))
                if _MEDICIONES is not None:
                    registro["camaras"] = _filas_exportadas(r, formato, disp)
                cuerpo = _xlsx_crudo(r.content) if formato == "xlsx" and "xlsx" in disp else r.content.decode("utf-8")
                resultados[f"export_{estado}_{formato}{variante}"] = (
                    f"status={r.status_code}\ncontent-type={r.headers.get('content-type')}\n"
                    f"content-disposition={disp}\n---\n{cuerpo}"
                )


def _correr(solo: str | None) -> dict[str, str]:
    """Devuelve {nombre: contenido_texto}; `solo` limita a 'smart' o 'export'."""
    from fastapi.testclient import TestClient

    from web.app.main import app

    resultados: dict[str, str] = {}
    with _sesion_linea_base(app):
        cliente = TestClient(app)
        if solo in (None, "smart"):
            _correr_smart(cliente, resultados)
        if solo in (None, "export"):
            _correr_export(cliente, resultados)
    return resultados


@contextlib.contextmanager
def _sesion_linea_base(app):
    """Sesión admin simulada que sirve para las dos versiones del router de Cámaras: la que toma la
    sesión con `Depends(usuario_sesion)` (override en `app.dependency_overrides`) y la anterior, que
    llamaba a `_require_auth(request)` importado en el router (parche sobre ese nombre). Así se puede
    grabar la línea base con el código anterior y comparar con el nuevo usando este mismo script."""
    from web.app.routers import camaras

    with contextlib.ExitStack() as pila:
        if hasattr(camaras, "_require_auth"):
            pila.enter_context(patch.object(camaras, "_require_auth", return_value=("linea-base", "admin")))
        try:
            from web.app.dependencias import UsuarioSesion, usuario_admin, usuario_sesion
        except ImportError:
            pass
        else:
            usuario = UsuarioSesion("linea-base", "admin")
            previos = {dep: app.dependency_overrides.get(dep) for dep in (usuario_sesion, usuario_admin)}
            for dep in previos:
                app.dependency_overrides[dep] = lambda: usuario

            def _restaurar() -> None:
                for dep, previo in previos.items():
                    if previo is None:
                        app.dependency_overrides.pop(dep, None)
                    else:
                        app.dependency_overrides[dep] = previo

            pila.callback(_restaurar)
        yield


def _grabar(salida: Path, solo: str | None) -> int:
    salida.mkdir(parents=True, exist_ok=True)
    resultados = _correr(solo)
    for nombre, contenido in resultados.items():
        with open(salida / f"{nombre}.json", "w", encoding="utf-8", newline="") as f:
            f.write(contenido)
    print(f"Grabadas {len([n for n in resultados if n[0] != '_'])} consultas en {salida}")
    return 0


def _comparar(salida: Path, solo: str | None) -> int:
    if not salida.is_dir():
        _abortar(f"no existe el directorio de línea base {salida}")
    resultados = _correr(solo)
    prefijos = (f"{solo}_",) if solo else ("smart_", "export_")
    grabados = {p.stem for p in salida.glob("*.json") if p.stem.startswith(prefijos)}
    actuales = {n for n in resultados if n.startswith(prefijos)}
    diferencias = 0
    for nombre in sorted(grabados | actuales):
        if nombre not in actuales or nombre not in grabados:
            print(f"DIFERENCIA {nombre}: presente sólo en {'línea base' if nombre in grabados else 'corrida actual'}")
            diferencias += 1
            continue
        with open(salida / f"{nombre}.json", encoding="utf-8", newline="") as f:
            antes = f.read()
        if antes != resultados[nombre]:
            diferencias += 1
            diff = difflib.unified_diff(
                antes.splitlines(keepends=True),
                resultados[nombre].splitlines(keepends=True),
                "linea_base",
                "actual",
                lineterm="",
                n=2,
            )
            print(f"DIFERENCIA {nombre}:")
            print("".join(repr(linea) + "\n" for linea in list(diff)[:60]))
    if diferencias:
        print(f"FALLA: {diferencias} diferencias", file=sys.stderr)
        return 1
    print(f"OK: {len(actuales)} consultas idénticas a la línea base")
    return 0


def _resumen_grupo(nombre: str, filas: list[dict]) -> None:
    if not filas:
        return
    seg = [f["segundos"] for f in filas]
    sent = [f["sentencias"] for f in filas]
    print(
        f"{nombre}: consultas={len(filas)} tiempo_total={sum(seg):.1f}s mediana={statistics.median(seg):.2f}s "
        f"peor={max(seg):.2f}s sentencias_total={sum(sent)} mediana={statistics.median(sent):.0f} peor={max(sent)}"
    )


def _medir(salida: Path, solo: str | None) -> int:
    global _MEDICIONES
    salida.mkdir(parents=True, exist_ok=True)
    _MEDICIONES = []
    try:
        _correr(solo)
        mediciones = _MEDICIONES
    finally:
        _MEDICIONES = None
    with open(salida / "tiempos.json", "w", encoding="utf-8") as f:
        json.dump(mediciones, f, ensure_ascii=False, indent=1)
    smart = [m for m in mediciones if m["nombre"].startswith("smart_")]
    _resumen_grupo("smart (todas)", smart)
    _resumen_grupo("smart con terminos", [m for m in smart if not m["nombre"].startswith("smart_vacio_")])
    _resumen_grupo("smart sin terminos", [m for m in smart if m["nombre"].startswith("smart_vacio_")])
    _resumen_grupo("export", [m for m in mediciones if m["nombre"].startswith("export_")])
    print(f"Escrito {salida / 'tiempos.json'}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("accion", choices=["grabar", "comparar", "medir"])
    parser.add_argument("--salida", required=True, type=Path)
    parser.add_argument("--solo", choices=["smart", "export"], default=None, help="limita a un grupo de consultas")
    args = parser.parse_args()
    _verificar_base_dev()
    accion = {"grabar": _grabar, "comparar": _comparar, "medir": _medir}[args.accion]
    return accion(args.salida, args.solo)


if __name__ == "__main__":
    sys.exit(main())
