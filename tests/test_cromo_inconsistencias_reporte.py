# Nombre de archivo: test_cromo_inconsistencias_reporte.py
# Ubicación de archivo: tests/test_cromo_inconsistencias_reporte.py
# Descripción: Pruebas del reporte XLSX de inconsistencias de Cromo (texto "qué falta", hojas por
# localidad), del aviso de casos nuevos (WS + Slack, best-effort) y del endpoint de exportación

from __future__ import annotations

import io
from types import SimpleNamespace
from typing import Any

import pytest
from openpyxl import load_workbook

from core.services.cromo import inconsistencias_cromo as inc
from core.services.cromo import inconsistencias_reporte as rep
from tests.test_web_cromo_camino_optico import _cliente_admin, _cliente_user, _fake_session

URL = "/api/admin/infra/inconsistencias-cromo/exportar.xlsx"


# ── Texto y hojas ────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "tipo, evidencia, esperado",
    [
        ("CABLE_SIN_EXTREMO", {"cable": {"extremo_a": 5, "extremo_b": None}}, "El cable no tiene extremo B en Cromo"),
        ("CABLE_SIN_EXTREMO", {"cable": {"extremo_a": None, "extremo_b": 5}}, "El cable no tiene extremo A en Cromo"),
        ("CABLE_SIN_EXTREMO", {"cable": {}}, "El cable no tiene extremo A ni extremo B en Cromo"),
        ("ISLA_FUSION", {"contenedor": {"es_odf": False}}, "Fusión aislada: ningún pelo continúa"),
        ("ISLA_FUSION", {"contenedor": {"es_odf": True}}, "Fusión aislada: ningún pelo continúa (dentro de un ODF)"),
        ("ELEMENTO_ISLA", None, "Elemento sin continuidad en sus cables"),
    ],
)
def test_que_falta(tipo, evidencia, esperado):
    assert rep.que_falta(tipo, evidencia) == esperado


def test_nombre_hoja_saneado_y_unico():
    usadas: set[str] = set()
    assert rep.nombre_hoja(None, usadas) == "Sin localidad"
    assert rep.nombre_hoja("  ", usadas) == "Sin localidad (2)"
    sucio = rep.nombre_hoja("A/B:C*D?[E]\\F", usadas)
    assert not any(c in sucio for c in "[]:*?/\\")
    largo = "L" * 50
    h1 = rep.nombre_hoja(largo, usadas)
    h2 = rep.nombre_hoja(largo, usadas)
    assert len(h1) == 31 and len(h2) <= 31 and h1 != h2
    # Sin distinguir mayúsculas (Excel las considera iguales).
    assert rep.nombre_hoja("korn", usadas) != rep.nombre_hoja("KORN", usadas)
    assert rep.nombre_hoja("'abc'", set()) == "abc"


def _item(**kw) -> dict[str, Any]:
    base = {
        "tipo": "CABLE_SIN_EXTREMO",
        "revision": "MANUAL",
        "estado": "PENDIENTE",
        "nombre": "F-TEST-1",
        "objeto_n_id": 123,
        "localidad": "Korn",
        "evidencia": {"cable": {"extremo_a": 1, "extremo_b": None}, "servicios": ["100", "200"]},
        "primera_vez": None,
        "ultima_vez": None,
    }
    base.update(kw)
    return base


def test_generar_xlsx_una_hoja_por_localidad():
    filas = [
        rep.fila_reporte(_item()),
        rep.fila_reporte(_item(objeto_n_id=124, localidad="Korn")),
        rep.fila_reporte(_item(objeto_n_id=125, localidad="La Plata")),
        rep.fila_reporte(_item(objeto_n_id=126, localidad=None)),
    ]
    wb = load_workbook(io.BytesIO(rep.generar_xlsx(filas)))
    assert sorted(wb.sheetnames) == ["Korn", "La Plata", "Sin localidad"]
    hoja = wb["Korn"]
    assert [c.value for c in hoja[1]] == list(rep.COLUMNAS)
    assert hoja.max_row == 3
    fila = {rep.COLUMNAS[i]: c.value for i, c in enumerate(hoja[2])}
    assert fila["ID Cromo"] == 123
    assert fila["Qué falta"] == "El cable no tiene extremo B en Cromo"
    assert fila["Servicios afectados"] == "100, 200"
    assert fila["Tipo"] == "Cable sin extremo"


def test_generar_xlsx_vacio_tiene_encabezados():
    wb = load_workbook(io.BytesIO(rep.generar_xlsx([])))
    assert wb.sheetnames == ["Sin localidad"]
    assert [c.value for c in wb.active[1]] == list(rep.COLUMNAS)


# ── Aviso de casos nuevos ────────────────────────────────────────────────────


@pytest.fixture
def sin_testing(monkeypatch):
    monkeypatch.delenv("TESTING", raising=False)


@pytest.fixture
def aviso(monkeypatch):
    """Instrumenta WS, config y envío a Slack; devuelve el registro de llamadas."""
    reg: dict[str, Any] = {"ws": [], "slack": [], "config": ([], False), "token": "xoxb-test"}

    async def _ws(payload):
        reg["ws"].append(dict(payload))

    async def _config(_sesion):
        return reg["config"]

    async def _filas(_sesion, claves):
        return [rep.fila_reporte(_item())]

    def _slack(canales, token, texto, xlsx, archivo):
        reg["slack"].append((list(canales), token, texto, xlsx, archivo))
        return len(canales)

    monkeypatch.setattr(rep, "_publicar_ws", _ws)
    monkeypatch.setattr(rep, "_leer_config_slack", _config)
    monkeypatch.setattr(rep, "filas_de_casos", _filas)
    monkeypatch.setattr(rep, "_enviar_slack_sync", _slack)
    monkeypatch.setattr(rep, "_token_slack", lambda: reg["token"])
    return reg


def _resumen(nuevos=3, auto=1, previos=10, claves=None):
    return SimpleNamespace(
        nuevos=nuevos,
        cuarentena_auto=auto,
        casos_previos=previos,
        nuevos_claves=claves if claves is not None else [("ISLA_FUSION", 1)],
    )


@pytest.mark.asyncio
async def test_sin_nuevos_no_avisa_nada(aviso):
    aviso["config"] = (["C1"], True)
    assert await rep.avisar_casos_nuevos(object(), _resumen(0, 0)) == {"ws": False, "slack": False}
    assert aviso["ws"] == [] and aviso["slack"] == []


@pytest.mark.asyncio
async def test_slack_inactivo_o_sin_canales_solo_ws(aviso):
    for config in (([], False), (["C1"], False), ([], True)):
        aviso["config"] = config
        assert await rep.avisar_casos_nuevos(object(), _resumen()) == {"ws": True, "slack": False}
    assert len(aviso["ws"]) == 3 and aviso["slack"] == []
    assert aviso["ws"][0]["type"] == "cromo_inconsistencias_nuevas"
    assert aviso["ws"][0]["nuevos"] == 3 and aviso["ws"][0]["cuarentena_auto"] == 1
    assert "at" in aviso["ws"][0]


@pytest.mark.asyncio
async def test_slack_activo_envia_xlsx_a_los_canales(aviso):
    aviso["config"] = (["C1", "C2"], True)
    assert await rep.avisar_casos_nuevos(object(), _resumen(nuevos=1, auto=0)) == {"ws": True, "slack": True}
    canales, token, texto, xlsx, archivo = aviso["slack"][0]
    assert canales == ["C1", "C2"] and token == "xoxb-test"
    assert "*1*" in texto and archivo.endswith(".xlsx")
    assert load_workbook(io.BytesIO(xlsx)).sheetnames == ["Korn"]


@pytest.mark.asyncio
async def test_primera_deteccion_solo_resumen_sin_adjunto(aviso):
    aviso["config"] = (["C1"], True)
    res = await rep.avisar_casos_nuevos(object(), _resumen(nuevos=5000, auto=100, previos=0))
    assert res == {"ws": True, "slack": True}
    assert aviso["ws"][0]["nuevos"] == 5000
    _canales, _token, texto, xlsx, _archivo = aviso["slack"][0]
    assert xlsx is None
    assert "*5000*" in texto and "primera detección" in texto and "viewer" in texto


@pytest.mark.asyncio
async def test_adjunto_cortado_al_tope_con_nota(aviso, monkeypatch):
    aviso["config"] = (["C1"], True)
    monkeypatch.setenv("CROMO_INCONSISTENCIAS_MAX_FILAS_AVISO", "2")
    pedidas: list[int] = []

    async def _filas(_sesion, claves):
        pedidas.append(len(claves))
        return [rep.fila_reporte(_item(objeto_n_id=i)) for i, _ in enumerate(claves)]

    monkeypatch.setattr(rep, "filas_de_casos", _filas)
    claves = [("ISLA_FUSION", i) for i in range(5)]
    await rep.avisar_casos_nuevos(object(), _resumen(nuevos=5, auto=0, claves=claves))
    assert pedidas == [2]
    _c, _t, texto, xlsx, _a = aviso["slack"][0]
    assert "y 3 más" in texto
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb["Korn"].max_row == 3
    assert wb["Aviso"]["A2"].value == "y 3 más — ver el viewer"


def test_tope_de_filas_por_env(monkeypatch):
    monkeypatch.delenv("CROMO_INCONSISTENCIAS_MAX_FILAS_AVISO", raising=False)
    assert rep.max_filas_aviso() == 500
    monkeypatch.setenv("CROMO_INCONSISTENCIAS_MAX_FILAS_AVISO", "x")
    assert rep.max_filas_aviso() == 500


def test_servicios_parciales_se_marcan():
    ev = {"servicios": [str(i) for i in range(20)]}
    f = rep.fila_reporte(_item(evidencia=ev, servicios=25))
    assert f["Servicios afectados"].endswith("(+5)")
    f = rep.fila_reporte(_item(evidencia=ev))
    assert "≥20" in f["Servicios afectados"]
    f = rep.fila_reporte(_item(evidencia={"servicios": ["1"]}, servicios=1))
    assert f["Servicios afectados"] == "1"


@pytest.mark.asyncio
async def test_slack_sin_token_se_omite(aviso):
    aviso["config"] = (["C1"], True)
    aviso["token"] = ""
    assert await rep.avisar_casos_nuevos(object(), _resumen()) == {"ws": True, "slack": False}
    assert aviso["slack"] == []


@pytest.mark.asyncio
async def test_best_effort_nunca_levanta(aviso, monkeypatch):
    async def _roto(*_a, **_k):
        raise RuntimeError("redis caído")

    aviso["config"] = (["C1"], True)
    monkeypatch.setattr(rep, "_publicar_ws", _roto)
    # WS caído no impide el aviso a Slack.
    assert await rep.avisar_casos_nuevos(object(), _resumen()) == {"ws": False, "slack": True}
    monkeypatch.setattr(rep, "_leer_config_slack", _roto)
    assert await rep.avisar_casos_nuevos(object(), _resumen()) == {"ws": False, "slack": False}
    # Resumen inválido tampoco levanta.
    assert await rep.avisar_casos_nuevos(object(), object()) == {"ws": False, "slack": False}


@pytest.mark.asyncio
async def test_ejecutar_corrida_llama_al_aviso_solo_si_ok_y_aplicado(monkeypatch):
    llamadas: list[Any] = []

    async def _aviso(sesion, resumen, *, corrida_id=None):
        llamadas.append((resumen, corrida_id))

    monkeypatch.setattr(inc, "avisar_casos_nuevos", _aviso)

    class _Corrida:
        total_objetivo = None
        estado = None
        finalizada_at = None
        leidas = 0
        creadas = actualizadas = errores = 0

    class _Sesion:
        async def get(self, *_a):
            return _Corrida()

        async def commit(self):
            pass

        async def rollback(self):
            pass

    async def _evento(*_a, **_k):
        pass

    monkeypatch.setattr(inc.ingesta, "registrar_evento", _evento)
    resumen = inc.ResumenDeteccion(
        por_tipo_revision={}, total=0, confirmaciones=inc.ResumenConfirmaciones(), aplicado=True, nuevos=2
    )

    async def _detectar(*_a, **_k):
        return resumen

    monkeypatch.setattr(inc, "detectar", _detectar)
    assert await inc._ejecutar_corrida(None, _Sesion(), 9) == "OK"
    assert llamadas == [(resumen, 9)]

    async def _falla(*_a, **_k):
        raise RuntimeError("x")

    llamadas.clear()
    monkeypatch.setattr(inc, "detectar", _falla)
    assert await inc._ejecutar_corrida(None, _Sesion(), 9) == "FALLIDA"
    assert llamadas == []


@pytest.mark.asyncio
async def test_ejecutar_corrida_ok_pero_no_aplicado_no_avisa(monkeypatch):
    llamadas: list[Any] = []

    async def _aviso(*_a, **_k):
        llamadas.append(1)

    monkeypatch.setattr(inc, "avisar_casos_nuevos", _aviso)

    class _Corrida:
        total_objetivo = None
        estado = None
        finalizada_at = None
        leidas = 0
        creadas = actualizadas = errores = 0

    class _Sesion:
        async def get(self, *_a):
            return _Corrida()

        async def commit(self):
            pass

        async def rollback(self):
            pass

    async def _evento(*_a, **_k):
        pass

    async def _detectar(*_a, **_k):
        return inc.ResumenDeteccion(
            por_tipo_revision={}, total=0, confirmaciones=inc.ResumenConfirmaciones(), aplicado=False, nuevos=2
        )

    monkeypatch.setattr(inc.ingesta, "registrar_evento", _evento)
    monkeypatch.setattr(inc, "detectar", _detectar)
    assert await inc._ejecutar_corrida(None, _Sesion(), 9) == "OK"
    assert llamadas == []


# ── Endpoint ─────────────────────────────────────────────────────────────────


def test_exportar_solo_admin(monkeypatch):
    _fake_session(monkeypatch)
    assert _cliente_user(monkeypatch).get(URL).status_code == 403


def test_exportar_valida_filtros(monkeypatch):
    _fake_session(monkeypatch)
    client = _cliente_admin(monkeypatch)
    assert client.get(f"{URL}?tipo=OTRO").status_code == 400
    assert client.get(f"{URL}?estado=OTRO").status_code == 400
    assert client.get(f"{URL}?revision=OTRA").status_code == 400


def test_exportar_pasa_filtros_y_devuelve_xlsx(monkeypatch):
    _fake_session(monkeypatch)
    recibido: dict[str, Any] = {}

    async def _filas(_sesion, **kw):
        recibido.update(kw)
        return [rep.fila_reporte(_item()), rep.fila_reporte(_item(localidad="La Plata"))]

    monkeypatch.setattr(rep, "filas_reporte", _filas)
    client = _cliente_admin(monkeypatch)
    res = client.get(f"{URL}?tipo=CABLE_SIN_EXTREMO&estado=PENDIENTE&revision=MANUAL&localidad=Korn&servicio=1&servicio_id=7")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "inconsistencias_cromo_" in res.headers["content-disposition"]
    assert sorted(load_workbook(io.BytesIO(res.content)).sheetnames) == ["Korn", "La Plata"]
    assert recibido == {
        "tipo": "CABLE_SIN_EXTREMO",
        "estado": "PENDIENTE",
        "revision": "MANUAL",
        "localidad": "Korn",
        "servicio": "1",
        "servicio_id": 7,
    }
    # Sin filtros: todo None (el default de estados accionables lo aplica filas_reporte).
    recibido.clear()
    assert client.get(URL).status_code == 200
    assert set(recibido.values()) == {None}


# ── Config del aviso a Slack ─────────────────────────────────────────────────

URL_AVISO = "/api/admin/infra/inconsistencias-cromo/aviso-slack"


def test_aviso_slack_solo_admin(monkeypatch):
    _fake_session(monkeypatch)
    client = _cliente_user(monkeypatch)
    assert client.get(URL_AVISO).status_code == 403
    assert client.post(URL_AVISO, json={"slack_channels": "#a", "activo": True}).status_code == 403


def test_aviso_slack_get(monkeypatch):
    _fake_session(monkeypatch)

    async def _leer(_s):
        return {"slack_channels": "#a", "activo": True, "existe": True}

    monkeypatch.setattr(rep, "leer_config_aviso", _leer)
    res = _cliente_admin(monkeypatch).get(URL_AVISO)
    assert res.status_code == 200 and res.json()["slack_channels"] == "#a"


def test_aviso_slack_post_valida_y_guarda(monkeypatch):
    _fake_session(monkeypatch)
    guardado: list[tuple] = []

    async def _guardar(_s, canales, activo):
        guardado.append((canales, activo))

    class _Sesion:
        async def commit(self):
            pass

    class _CM:
        async def __aenter__(self):
            return _Sesion()

        async def __aexit__(self, *a):
            return False

    monkeypatch.setattr("db.session.AsyncSessionLocal", lambda: _CM())
    monkeypatch.setattr(rep, "guardar_config_aviso", _guardar)
    monkeypatch.setenv("TESTING", "true")
    client = _cliente_admin(monkeypatch)

    ok = client.post(URL_AVISO, json={"slack_channels": " #cromo-alertas , C0123456789 ", "activo": True})
    assert ok.status_code == 200
    assert guardado == [("#cromo-alertas,C0123456789", True)]
    # formato inválido, activar sin canal, demasiado largo
    assert client.post(URL_AVISO, json={"slack_channels": "canal sin numeral", "activo": False}).status_code == 400
    assert client.post(URL_AVISO, json={"slack_channels": "", "activo": True}).status_code == 400
    assert client.post(URL_AVISO, json={"slack_channels": ",".join(["#a"] * 400), "activo": False}).status_code == 400
    # desactivar sin canales es válido
    assert client.post(URL_AVISO, json={"slack_channels": "", "activo": False}).status_code == 200
    assert len(guardado) == 2


def test_aviso_slack_post_sin_csrf_403(monkeypatch, sin_testing):
    _fake_session(monkeypatch)
    res = _cliente_admin(monkeypatch).post(
        URL_AVISO, json={"slack_channels": "#a", "activo": False, "csrf_token": "otro"}
    )
    assert res.status_code == 403
