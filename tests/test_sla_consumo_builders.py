# Nombre de archivo: test_sla_consumo_builders.py
# Ubicación de archivo: tests/test_sla_consumo_builders.py
# Descripción: Gráficos, XLSX y DOCX del informe SLA consumido

from __future__ import annotations

import openpyxl
import pytest
from docx import Document

from core.sla_consumo import charts
from core.sla_consumo.docx_builder import construir_docx
from core.sla_consumo.xlsx_builder import HOJAS, construir_xlsx
from tests.test_sla_consumo_engine import resultado  # noqa: F401  (fixture reutilizada)


def test_charts_generan_png(resultado, tmp_path):
    for archivo in (charts.pareto_eventos(resultado, tmp_path / "a.png"),
                    charts.top_servicios(resultado, tmp_path / "b.png"),
                    charts.horas_por_grupo(resultado, tmp_path / "c.png"),
                    charts.servicio(resultado, "L1", tmp_path / "d.png")):
        assert archivo.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"


@pytest.mark.filterwarnings("ignore:datetime.datetime.utcnow:DeprecationWarning")
def test_xlsx_tiene_todas_las_hojas(resultado, tmp_path):
    wb = openpyxl.load_workbook(construir_xlsx(resultado, tmp_path / "x.xlsx"))
    assert wb.sheetnames == list(HOJAS)
    assert wb["Servicios"].max_row == len(resultado.servicios) + 1


def test_docx_con_graficos(resultado, tmp_path):
    doc = Document(construir_docx(resultado, tmp_path / "x.docx", top_servicios=2))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "SLA consumido" in texto and "2026-10-01" in texto
    assert len(doc.inline_shapes) >= 3 + 2   # 3 generales + 1 por servicio del top


def _con_inf(resultado):
    s = resultado.servicios.copy()
    s["pct_presupuesto"] = s["pct_presupuesto"].astype(float)
    s.loc[s.index[0], "pct_presupuesto"] = float("inf")
    s.loc[s.index[1], "pct_presupuesto"] = float("nan")
    resultado.servicios = s
    return resultado


@pytest.mark.filterwarnings("ignore:datetime.datetime.utcnow:DeprecationWarning")
def test_xlsx_sin_inf_ni_fallo(resultado, tmp_path):
    wb = openpyxl.load_workbook(construir_xlsx(_con_inf(resultado), tmp_path / "x.xlsx"))
    valores = [c.value for fila in wb["Servicios"].iter_rows() for c in fila]
    assert not any(isinstance(v, str) and "inf" in v.lower() for v in valores)
    assert not any(isinstance(v, float) and v == float("inf") for v in valores)


def test_docx_y_charts_con_inf_nan(resultado, tmp_path):
    res = _con_inf(resultado)
    assert charts.top_servicios(res, tmp_path / "t.png").exists()
    doc = Document(construir_docx(res, tmp_path / "x.docx", top_servicios=5))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "inf" not in texto.lower().replace("informe", "").replace("información", "")
    assert "nan" not in texto.lower()
    celdas = " ".join(c.text for t in doc.tables for f in t.rows for c in f.cells).lower()
    assert "inf" not in celdas and "nan" not in celdas


@pytest.mark.filterwarnings("ignore:datetime.datetime.utcnow:DeprecationWarning")
def test_xlsx_quita_tz(resultado, tmp_path):
    wb = openpyxl.load_workbook(construir_xlsx(resultado, tmp_path / "x.xlsx"))
    assert wb["Eventos"].max_row == len(resultado.eventos) + 1
