# Nombre de archivo: test_sla_consumo_docx.py
# Ubicación de archivo: tests/test_sla_consumo_docx.py
# Descripción: DOCX ejecutivo y exhaustivo del informe SLA consumido — fichas, encabezados repetidos, colores y top N

from __future__ import annotations

import re

import openpyxl
import pytest
from docx import Document
from docx.oxml.ns import qn

from core.sla_consumo import metricas as m
from core.sla_consumo import presentacion
from core.sla_consumo.docx_builder import (TOP_EVENTOS_DEFAULT, construir_docx_ejecutivo, construir_docx_exhaustivo,
                                           top_eventos_config)
from core.sla_consumo.presentacion import COLUMNAS
from core.sla_consumo.xlsx_builder import construir_xlsx
from tests.test_sla_consumo_engine import CARRIER, FO, OTROS, _calc, _r, _s

sin_utcnow = pytest.mark.filterwarnings("ignore:datetime.datetime.utcnow:DeprecationWarning")


@pytest.fixture
def res():
    servicios = [_s("100", cliente="ALFA"), _s("200", sla=None, cliente="BETA"), _s("300", cliente="GAMMA"),
                 _s("400", cliente="DELTA"), _s("500", cliente="EPSILON")]
    reclamos = [_r("1", "100", 27, FO, evento="E1", dia=1), _r("2", "100", 1, CARRIER, dia=2),
                _r("3", "300", 4, OTROS, evento="E1", dia=3), _r("4", "200", 2, FO, evento="E2", dia=4),
                _r("5", "999", 3, FO, dia=5), _r("6", "500", 26, FO, evento="E3", dia=6),
                _r("7", "300", 1, FO, evento="E4", dia=7), _r("8", "300", 0.5, FO, evento="E5", dia=8)]
    return _calc(servicios, reclamos)


def _titulos_ficha(doc, res) -> set[str]:
    fichas = {f"{s} — {c}" for s, c in zip(res.servicios["servicio_id"], res.servicios["nombre_cliente"])}
    return {p.text for p in doc.paragraphs if p.style.name == "Heading 2" and p.text in fichas}


def _textos_celdas(doc) -> list[str]:
    return [t.text or "" for t in doc.element.body.iter(qn("w:t"))]


def _tabla_con_encabezado(doc, etiqueta):
    return [t for t in doc.tables if etiqueta in [c.text for c in t.rows[0].cells]]


def test_ejecutivo_fichas_solo_agotados_y_excedidos(res, tmp_path):
    doc = Document(construir_docx_ejecutivo(res, tmp_path / "e.docx"))
    estados = dict(zip(res.servicios["servicio_id"], res.servicios["estado"]))
    assert estados["100"] == m.ESTADO_EXCEDIDO and estados["500"] == m.ESTADO_AGOTADO
    assert _titulos_ficha(doc, res) == {"100 — ALFA", "500 — EPSILON"}


def test_exhaustivo_fichas_de_todo_servicio_con_reclamos_y_tabla_sin_reclamos(res, tmp_path):
    doc = Document(construir_docx_exhaustivo(res, tmp_path / "x.docx"))
    assert _titulos_ficha(doc, res) == {"100 — ALFA", "200 — BETA", "300 — GAMMA", "500 — EPSILON"}
    encabezados = [COLUMNAS[c][0] for c in ("servicio_id", "nombre_cliente", "tipo_servicio", "sla_prometido",
                                            "presupuesto_h", "estado")]
    compacta = [t for t in doc.tables if [c.text for c in t.rows[0].cells] == encabezados]
    assert len(compacta) == 1
    assert [f.cells[0].text for f in compacta[0].rows[1:]] == ["400"]


@pytest.mark.parametrize("construir", [construir_docx_ejecutivo, construir_docx_exhaustivo])
def test_toda_tabla_repite_encabezado(res, tmp_path, construir):
    doc = Document(construir(res, tmp_path / "d.docx"))
    assert doc.tables
    for tabla in doc.tables:
        tr_pr = tabla.rows[0]._tr.trPr
        assert tr_pr is not None and tr_pr.find(qn("w:tblHeader")) is not None


def test_celda_semaforo_con_texto_y_sombreado(res, tmp_path):
    doc = Document(construir_docx_ejecutivo(res, tmp_path / "e.docx"))
    celdas = [c for t in doc.tables for f in t.rows for c in f.cells if c.text == m.AMARILLO]
    assert celdas
    fills = {shd.get(qn("w:fill")) for c in celdas for shd in c._tc.iter(qn("w:shd"))}
    assert "F1C40F" in fills


def test_celda_estado_coloreada(res, tmp_path):
    doc = Document(construir_docx_ejecutivo(res, tmp_path / "e.docx"))
    celdas = [c for t in doc.tables for f in t.rows for c in f.cells if c.text == m.ESTADO_EXCEDIDO]
    fills = {shd.get(qn("w:fill")) for c in celdas for shd in c._tc.iter(qn("w:shd"))}
    assert presentacion.COLORES_ESTADO[m.ESTADO_EXCEDIDO] in fills


@pytest.mark.parametrize("valor, esperado", [("3", 3), ("0", TOP_EVENTOS_DEFAULT), ("-2", TOP_EVENTOS_DEFAULT),
                                             ("abc", TOP_EVENTOS_DEFAULT), ("", TOP_EVENTOS_DEFAULT),
                                             (None, TOP_EVENTOS_DEFAULT)])
def test_top_eventos_config(monkeypatch, valor, esperado):
    if valor is None:
        monkeypatch.delenv("SLA_CONSUMO_TOP_EVENTOS", raising=False)
    else:
        monkeypatch.setenv("SLA_CONSUMO_TOP_EVENTOS", valor)
    assert TOP_EVENTOS_DEFAULT == 20
    assert top_eventos_config() == esperado


def test_ejecutivo_muestra_a_lo_sumo_n_eventos(res, tmp_path, monkeypatch):
    assert len(res.eventos) == 5
    monkeypatch.setenv("SLA_CONSUMO_TOP_EVENTOS", "3")
    doc = Document(construir_docx_ejecutivo(res, tmp_path / "e.docx"))
    principales = _tabla_con_encabezado(doc, "Rango observado")
    assert len(principales) == 1 and len(principales[0].rows) - 1 == 3
    assert "SLA_CONSUMO_TOP_EVENTOS" in "\n".join(p.text for p in doc.paragraphs)
    doc2 = Document(construir_docx_ejecutivo(res, tmp_path / "e2.docx", top_eventos=2))
    assert len(_tabla_con_encabezado(doc2, "Rango observado")[0].rows) - 1 == 2
    completo = Document(construir_docx_exhaustivo(res, tmp_path / "x.docx"))
    assert len(_tabla_con_encabezado(completo, "Rango observado")[0].rows) - 1 == 5


@pytest.mark.parametrize("construir", [construir_docx_ejecutivo, construir_docx_exhaustivo])
def test_sin_consumo_reemplaza_torta(tmp_path, construir):
    vacio = _calc([_s("100")], [_r("1", "100", 0, FO)])
    destino = construir(vacio, tmp_path / "v.docx")
    doc = Document(destino)
    assert "Sin consumo" in [p.text for p in doc.paragraphs]
    assert not doc.inline_shapes


def test_torta_y_magnitud_en_carpeta_graficos(res, tmp_path):
    destino = construir_docx_ejecutivo(res, tmp_path / "e.docx")
    carpeta = tmp_path / "e_graficos"
    assert carpeta.is_dir() and len(list(carpeta.glob("*.png"))) == 2
    assert len(Document(destino).inline_shapes) == 2


@pytest.mark.parametrize("construir", [construir_docx_ejecutivo, construir_docx_exhaustivo])
def test_ninguna_celda_inf_ni_nan(res, tmp_path, construir):
    s = res.servicios.copy()
    s.loc[s["servicio_id"] == "100", "pct_real"] = float("inf")
    s.loc[s["servicio_id"] == "500", "horas_restantes"] = float("-inf")
    res.servicios = s
    doc = Document(construir(res, tmp_path / "d.docx"))
    for texto in _textos_celdas(doc):
        assert not re.search(r"\b-?(inf|nan)\b", texto, re.IGNORECASE), texto


def test_leyendas_y_resumen(res, tmp_path):
    doc = Document(construir_docx_ejecutivo(res, tmp_path / "e.docx"))
    textos = "\n".join(_textos_celdas(doc))
    for concepto, _ in presentacion.LEYENDAS:
        assert concepto in textos
    assert presentacion.TOTALES["eventos"][0] in textos
    assert "Eventos + reclamos sin evento" not in textos


def test_pagina_a4_apaisada(res, tmp_path):
    doc = Document(construir_docx_ejecutivo(res, tmp_path / "e.docx"))
    seccion = doc.sections[0]
    assert seccion.page_width > seccion.page_height
    assert abs(seccion.page_width.cm - 29.7) < 0.1


def test_fmt_conserva_texto_no_numerico():
    for texto in ("Infinity", "nan", "1e400", "inf"):
        assert presentacion.fmt(texto, "texto") == texto
        assert not presentacion._faltante(texto)
    assert presentacion._faltante(float("inf")) and presentacion._faltante(float("nan"))
    assert presentacion.fmt(float("nan"), "horas") == "—" and presentacion.fmt(27.5, "horas") == "27:30"
    assert presentacion.fmt(0.997, "pct", "sla_entregado_oficial") == "99,70 %"


@sin_utcnow
def test_xlsx_y_docx_conservan_texto_infinity(tmp_path):
    res = _calc([_s("100", cliente="Infinity")], [_r("1", "100", 30, FO)])
    ws = openpyxl.load_workbook(construir_xlsx(res, tmp_path / "x.xlsx"))["Servicios"]
    j = [c.value for c in ws[1]].index(COLUMNAS["nombre_cliente"][0]) + 1
    assert ws.cell(2, j).value == "Infinity"
    doc = Document(construir_docx_ejecutivo(res, tmp_path / "e.docx"))
    assert _titulos_ficha(doc, res) == {"100 — Infinity"}


def test_xlsx_builder_reusa_faltante_de_presentacion():
    from core.sla_consumo import xlsx_builder
    assert xlsx_builder._faltante is presentacion._faltante


@sin_utcnow
@pytest.mark.parametrize("construir", [construir_docx_ejecutivo, construir_docx_exhaustivo])
def test_docx_no_muestra_primer_servicio(tmp_path, construir):
    servicios = [_s("200", primer="98765", sla=99.0), _s("300")]
    reclamos = [_r("1", "200", 90, FO, evento="E1", dia=1, primer="98765"), _r("2", "300", 2, FO, dia=2),
                _r("3", "777", 1, FO, dia=3, primer="55555")]
    r = _calc(servicios, reclamos)
    doc = Document(str(construir(r, tmp_path / "d.docx")))
    textos = _textos_celdas(doc)
    assert "98765" not in textos and not any("98765" in t for t in textos)
    assert not any(COLUMNAS["numero_primer_servicio"][0] in t or COLUMNAS["clave_servicio"][0] in t for t in textos)
    assert any(t.startswith("200 — ") for t in textos)
