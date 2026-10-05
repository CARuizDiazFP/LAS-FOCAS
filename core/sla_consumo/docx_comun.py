# Nombre de archivo: docx_comun.py
# Ubicación de archivo: core/sla_consumo/docx_comun.py
# Descripción: Piezas compartidas de los DOCX del informe SLA consumido — página, tablas rápidas, resumen y ficha de servicio

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from core.sla_consumo import charts
from core.sla_consumo import metricas as m
from core.sla_consumo.presentacion import (COLORES_ESTADO, COLORES_SEMAFORO, COLORES_SUFICIENTE, COLUMNAS, LEYENDAS,
                                           TOTALES, VACIO, fmt, fmt_bool, fmt_fecha, fmt_horas)
from core.utils.excel_duraciones import TZ_AR

ESTILO_TABLA = "Table Grid"
_ESTILO_CELDA = "Celda"
_ESTILO_ENCABEZADO = "CeldaEncabezado"
_ESTILO_SUBTITULO = "SubtituloFicha"
_MARGEN_CELDA_DXA = 57                  # ≈ 0,1 cm a cada lado: más texto útil en columnas angostas
_FONDO_ENCABEZADO = "1F3864"
_CLAROS = {"F1C40F", "BDC3C7"}          # fondos donde el texto blanco no se lee
_ILEGALES = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")
_MARGEN = Cm(1.27)
_ALTO_IMAGEN_MAX_CM = 16.5
SIN_EVENTO = "Sin evento"
SIN_CONSUMO = "Sin consumo"

_COLORES = {"semaforo": COLORES_SEMAFORO, "indicador": COLORES_SEMAFORO, "estado": COLORES_ESTADO,
            "estado_servicio": COLORES_ESTADO, "suficiente_solo": COLORES_SUFICIENTE}


def etiqueta(columna: str) -> str:
    return COLUMNAS[columna][0]


def texto(valor: object, columna: str) -> str:
    """Valor formateado según `COLUMNAS`; nunca devuelve "nan"/"inf" para un número faltante."""
    return fmt(valor, COLUMNAS[columna][1], columna)


def color(columna: str, valor: str) -> str | None:
    paleta = _COLORES.get(columna)
    return paleta.get(valor) if paleta else None


# --------------------------------------------------------------------------------------------- documento

def nuevo_documento():
    doc = Document()
    seccion = doc.sections[0]
    seccion.orientation = WD_ORIENT.LANDSCAPE
    seccion.page_width, seccion.page_height = Cm(29.7), Cm(21.0)
    seccion.left_margin = seccion.right_margin = seccion.top_margin = seccion.bottom_margin = _MARGEN
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    celda = doc.styles.add_style(_ESTILO_CELDA, WD_STYLE_TYPE.PARAGRAPH)
    celda.base_style = normal
    celda.font.size = Pt(7)
    celda.paragraph_format.space_before = celda.paragraph_format.space_after = Pt(0)
    celda.paragraph_format.line_spacing = 1.0
    encabezado = doc.styles.add_style(_ESTILO_ENCABEZADO, WD_STYLE_TYPE.PARAGRAPH)
    encabezado.base_style = celda
    encabezado.font.bold = True
    encabezado.font.size = Pt(6.5)
    encabezado.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    subtitulo = doc.styles.add_style(_ESTILO_SUBTITULO, WD_STYLE_TYPE.PARAGRAPH)
    subtitulo.base_style = normal
    subtitulo.font.bold, subtitulo.font.size = True, Pt(9)
    subtitulo.paragraph_format.space_before, subtitulo.paragraph_format.space_after = Pt(6), Pt(2)
    subtitulo.paragraph_format.keep_with_next = True
    # python-docx resuelve sección y estilos recorriendo el cuerpo entero en cada add_*: con miles de fichas
    # eso es cuadrático. Se resuelven una vez y los bloques se insertan como XML antes del sectPr final.
    doc._sla = {
        "ancho_dxa": int((seccion.page_width - seccion.left_margin - seccion.right_margin) / 635),  # EMU → twips
        "sectPr": doc.element.body.sectPr,
        "estilos": {n: doc.styles[n].style_id for n in ("Title", "Heading 1", "Heading 2", ESTILO_TABLA,
                                                         _ESTILO_SUBTITULO)},
    }
    return doc


def _insertar(doc, elemento) -> None:
    doc._sla["sectPr"].addprevious(elemento)


def parrafo(doc, contenido: str, estilo: str | None = None) -> None:
    ppr = f'<w:pPr><w:pStyle w:val="{doc._sla["estilos"][estilo]}"/></w:pPr>' if estilo else ""
    run = f'<w:r><w:t xml:space="preserve">{_limpio(contenido)}</w:t></w:r>' if contenido else ""
    _insertar(doc, parse_xml(f"<w:p {nsdecls('w')}>{ppr}{run}</w:p>"))


def titulo(doc, contenido: str, nivel: int) -> None:
    parrafo(doc, contenido, "Title" if nivel == 0 else f"Heading {nivel}")


def _limpio(valor: str) -> str:
    return escape(_ILEGALES.sub("", valor))


def _parrafo_xml(contenido: str, estilo: str, rpr: str) -> str:
    lineas = contenido.split("\n")
    runs = "<w:br/>".join(f'<w:t xml:space="preserve">{_limpio(linea)}</w:t>' for linea in lineas)
    return f'<w:p><w:pPr><w:pStyle w:val="{estilo}"/></w:pPr><w:r>{rpr}{runs}</w:r></w:p>'


def _celda_xml(contenido: str, ancho: int, fondo: str | None, estilo: str) -> str:
    shd, rpr = "", ""
    if fondo:
        shd = f'<w:shd w:val="clear" w:color="auto" w:fill="{fondo}"/>'
        if estilo == _ESTILO_CELDA:
            rpr = f'<w:rPr><w:b/><w:color w:val="{"000000" if fondo in _CLAROS else "FFFFFF"}"/></w:rPr>'
    return (f'<w:tc><w:tcPr><w:tcW w:w="{ancho}" w:type="dxa"/>{shd}</w:tcPr>'
            f'{_parrafo_xml(contenido, estilo, rpr)}</w:tc>')


def tabla(doc, encabezados: list[str], filas: list[list[str]], fondos: list[list[str | None]] | None = None,
          anchos: list[float] | None = None, no_partir_filas: bool = True, fraccion: float = 1.0):
    """Tabla (estilo Table Grid, ancho fijo) con el encabezado repetido en cada página, en un solo parseo XML.

    `filas` llegan ya formateadas como texto; `fondos` (opcional) da el color hex de cada celda o None;
    `fraccion` es la parte del ancho útil de la página que ocupa la tabla.
    """
    k = len(encabezados)
    total = int(doc._sla["ancho_dxa"] * fraccion)
    relativos = anchos or [1.0] * k
    suma = sum(relativos)
    dxa = [int(total * r / suma) for r in relativos]
    tr_pr = "<w:trPr><w:cantSplit/></w:trPr>" if no_partir_filas else ""
    partes = [f"<w:tbl {nsdecls('w')}><w:tblPr><w:tblStyle w:val=\"{doc._sla['estilos'][ESTILO_TABLA]}\"/>",
              f'<w:tblW w:w="{sum(dxa)}" w:type="dxa"/><w:tblLayout w:type="fixed"/>'
              f'<w:tblCellMar><w:left w:w="{_MARGEN_CELDA_DXA}" w:type="dxa"/>'
              f'<w:right w:w="{_MARGEN_CELDA_DXA}" w:type="dxa"/></w:tblCellMar>',
              '<w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" '
              'w:noHBand="0" w:noVBand="1"/></w:tblPr><w:tblGrid>',
              *(f'<w:gridCol w:w="{a}"/>' for a in dxa), "</w:tblGrid>",
              "<w:tr><w:trPr><w:cantSplit/><w:tblHeader/></w:trPr>",
              *(_celda_xml(e, a, _FONDO_ENCABEZADO, _ESTILO_ENCABEZADO) for e, a in zip(encabezados, dxa)),
              "</w:tr>"]
    for i, fila in enumerate(filas):
        colores = fondos[i] if fondos else [None] * k
        partes.append(f"<w:tr>{tr_pr}")
        partes.extend(_celda_xml(v, a, c, _ESTILO_CELDA) for v, a, c in zip(fila, dxa, colores))
        partes.append("</w:tr>")
    partes.append("</w:tbl>")
    _insertar(doc, parse_xml("".join(partes)))


def tabla_df(doc, df: pd.DataFrame, columnas: list[str], anchos: list[float] | None = None,
             vacio: str = "Sin registros.", no_partir_filas: bool = True):
    """Tabla de un DataFrame con etiquetas y formatos de `COLUMNAS`; las columnas con paleta se colorean."""
    if df.empty:
        parrafo(doc, vacio)
        return
    textos = [[texto(v, c) for v in df[c]] for c in columnas]
    filas = [list(f) for f in zip(*textos)]
    fondos = None
    if any(c in _COLORES for c in columnas):
        por_col = [[color(c, v) for v in col] if c in _COLORES else [None] * len(df)
                   for c, col in zip(columnas, textos)]
        fondos = [list(f) for f in zip(*por_col)]
    tabla(doc, [etiqueta(c) for c in columnas], filas, fondos, anchos, no_partir_filas)


def imagen(doc, ruta: Path, ancho_cm: float) -> None:
    with Image.open(ruta) as img:
        w, h = img.size
    if ancho_cm * h / w > _ALTO_IMAGEN_MAX_CM:
        doc.add_picture(str(ruta), height=Cm(_ALTO_IMAGEN_MAX_CM))
    else:
        doc.add_picture(str(ruta), width=Cm(ancho_cm))


# --------------------------------------------------------------------------------------------- secciones comunes

def encabezado(doc, res, nombre: str) -> None:
    titulo(doc, nombre, 0)
    inicio = res.fecha_corte - dt.timedelta(days=365)
    parrafo(doc, f"Fecha de corte: {res.fecha_corte:%d/%m/%Y}")
    parrafo(doc, f"Ventana de análisis (365 días): {inicio:%d/%m/%Y} a {res.fecha_corte:%d/%m/%Y}")
    parrafo(doc, f"Fecha de generación: {dt.datetime.now(TZ_AR):%d/%m/%Y %H:%M}")


def resumen(doc, res) -> None:
    titulo(doc, "Resumen global", 1)
    parrafo(doc, "Reclamos y horas netas cubren todos los reclamos; el resto de las horas, sólo los "
                      "reclamos vinculados a un servicio del universo.")
    filas = [[eti, fmt(res.totales[clave], formato)] for clave, (eti, formato) in TOTALES.items()
             if clave in res.totales]
    tabla(doc, ["Concepto", "Valor"], filas, anchos=[3, 1], fraccion=0.6)


def leyendas(doc) -> None:
    titulo(doc, "Leyendas", 1)
    tabla(doc, ["Concepto", "Explicación"], [[c, e] for c, e in LEYENDAS], anchos=[1, 3])


def torta_global(doc, res, graficos: Path) -> None:
    titulo(doc, "Composición de las horas computables por causa", 1)
    ruta = charts.torta_causas(res, graficos / "torta_causas.png")
    if ruta is None:
        parrafo(doc, SIN_CONSUMO)
    else:
        imagen(doc, ruta, 16)
    t = res.totales
    parrafo(doc, f"{TOTALES['horas_excluidas'][0]}: {fmt_horas(t.get('horas_excluidas'))} h. "
                      f"{TOTALES['horas_no_vinculadas'][0]}: {fmt_horas(t.get('horas_no_vinculadas'))} h "
                      "(fuera del universo de servicios, no incluidas en la composición).")


# --------------------------------------------------------------------------------------------- eventos

_CAUSAS = ((m.CAUSA_FO_GENERAL, "horas_fo_general"), (m.CAUSA_FO_COD3, "horas_fo_cod3"),
           (m.CAUSA_CARRIER, "horas_carrier"), (m.CAUSA_OTROS, "horas_otros"))


def eventos_ordenados(res) -> pd.DataFrame:
    ev = res.eventos
    return ev.assign(_h=pd.to_numeric(ev["horas_servicio"], errors="coerce").fillna(0.0)) \
             .sort_values("_h", ascending=False, kind="mergesort").drop(columns="_h")


def _indicador_evento(n: object, lista: object) -> str:
    cantidad = fmt(n, "entero")
    lista_txt = "" if lista is None or (isinstance(lista, float) and pd.isna(lista)) else str(lista)
    return f"{cantidad}: {lista_txt}" if lista_txt else cantidad


def _causas_evento(fila) -> str:
    partes = []
    for causa, col in _CAUSAS:
        valor = pd.to_numeric(fila[col], errors="coerce")
        if pd.notna(valor) and valor > m.EPS_HORAS:
            partes.append(f"{causa}: {fmt_horas(valor)} h")
    return "\n".join(partes) or VACIO


def tabla_eventos(doc, eventos: pd.DataFrame) -> None:
    if eventos.empty:
        parrafo(doc, "Sin eventos reales.")
        return
    encabezados = [etiqueta("numero_evento"), "Rango observado", etiqueta("reclamos"),
                   etiqueta("servicios_afectados"), etiqueta("horas_servicio"),
                   etiqueta("servicios_agotados_o_excedidos"), etiqueta("servicios_suficiente_solo"),
                   etiqueta("servicios_cruza_umbral"), etiqueta("servicios_determinante"), "Composición por causa"]
    filas = []
    for _, f in eventos.iterrows():
        filas.append([texto(f["numero_evento"], "numero_evento"),
                      f"{fmt_fecha(f['inicio_observado'])} a\n{fmt_fecha(f['cierre_observado'])}",
                      texto(f["reclamos"], "reclamos"), texto(f["servicios_afectados"], "servicios_afectados"),
                      texto(f["horas_servicio"], "horas_servicio"),
                      texto(f["servicios_agotados_o_excedidos"], "servicios_agotados_o_excedidos"),
                      _indicador_evento(f["n_suficiente_solo"], f["servicios_suficiente_solo"]),
                      _indicador_evento(f["n_cruza_umbral"], f["servicios_cruza_umbral"]),
                      _indicador_evento(f["n_determinante"], f["servicios_determinante"]),
                      _causas_evento(f)])
    tabla(doc, encabezados, filas, anchos=[1.1, 1.5, 0.9, 1.0, 1.0, 1.1, 2.0, 2.0, 2.0, 1.9], no_partir_filas=False)
    parrafo(doc, "Los indicadores muestran la cantidad de servicios y, tras los dos puntos, sus IDs visibles. "
                      "Las horas-servicio no son la duración del evento y no deben compararse contra el "
                      "presupuesto de un único servicio: " + dict(LEYENDAS)["Horas-servicio ≠ duración del evento"])


# --------------------------------------------------------------------------------------------- ficha de servicio

_METRICAS_FICHA = ["tipo_servicio", "sla_prometido", "presupuesto_h", "horas_computables", "horas_excluidas",
                   "horas_fo_general", "horas_fo_cod3", "horas_carrier", "horas_otros", "pct_visible",
                   "horas_excedidas", "pct_real", "horas_restantes", "aporte_fo_pct", "participacion_fo_pct",
                   "reclamos", "eventos_reales", "reclamos_sin_evento", "semaforo", "estado"]

_RECLAMOS_FICHA = ["numero_reclamo", "_evento", "fecha_inicio", "fecha_cierre", "tipo_solucion", "_grupo_codigo",
                   "indicador", "horas_netas", "cuenta_sla", "horas_computables", "pct_aporte_real",
                   "consumo_acum_antes", "consumo_acum_despues", "restantes_acum", "excedidas_acum", "_alcanza"]
_RECLAMOS_ANCHOS = [1.3, 1.3, 1.6, 1.6, 2.3, 1.6, 1.5, 1.3, 1.15, 1.85, 1.95, 1.55, 1.55, 1.5, 1.5, 1.8]
_EVENTOS_FICHA = ["numero_evento", "reclamos_ids", "horas_computables", "pct_aporte_real", "suficiente_solo",
                  "cruza_umbral", "determinante_al_excluir", "participa_en_agotado_o_excedido"]
_EVENTOS_ANCHOS = [1.2, 2.4, 1.0, 1.0, 1.0, 1.0, 1.0, 1.2]
_ETIQUETAS_EXTRA = {"_evento": etiqueta("numero_evento"),
                    "_grupo_codigo": f"{etiqueta('grupo_cierre')} / {etiqueta('codigo_cierre')}",
                    "_alcanza": "Alcanza / supera el presupuesto por primera vez"}


def _alcanza(alcanza: object, supera: object) -> str:
    a, s = fmt_bool(alcanza) == "Sí", fmt_bool(supera) == "Sí"
    if a and s:
        return "Alcanza y supera"
    return "Alcanza" if a else "Supera" if s else "No"


class FichasServicio:
    """Pre-formatea una sola vez Servicio x Reclamo y Evento x Servicio para escribir miles de fichas rápido."""

    def __init__(self, res):
        sr = res.servicio_reclamo
        columnas = {c: [texto(v, c) for v in sr[c]] for c in _RECLAMOS_FICHA if not c.startswith("_")}
        for c in ("fecha_inicio", "fecha_cierre"):     # fecha y hora en dos líneas: la columna es angosta
            columnas[c] = [v.replace(" ", "\n", 1) for v in columnas[c]]
        columnas["_evento"] = [SIN_EVENTO if v == VACIO else v for v in (texto(x, "numero_evento")
                                                                         for x in sr["numero_evento"])]
        columnas["_grupo_codigo"] = [f"{texto(g, 'grupo_cierre')} / {texto(c, 'codigo_cierre')}"
                                     for g, c in zip(sr["grupo_cierre"], sr["codigo_cierre"])]
        columnas["_alcanza"] = [_alcanza(a, s) for a, s in zip(sr["alcanza_primera_vez"], sr["supera_primera_vez"])]
        self._reclamos = [list(f) for f in zip(*(columnas[c] for c in _RECLAMOS_FICHA))]
        j = _RECLAMOS_FICHA.index("indicador")
        self._fondos_reclamos = [[COLORES_SEMAFORO.get(f[j]) if i == j else None for i in range(len(f))]
                                 for f in self._reclamos]
        self._reclamos_por_clave = sr.groupby("clave_servicio", sort=False).indices
        self._eventos = res.evento_servicio
        self._eventos_por_clave = res.evento_servicio.groupby("clave_servicio", sort=False).indices

    def escribir(self, doc, fila_servicio) -> None:
        s = fila_servicio
        titulo(doc, f"{texto(s['servicio_id'], 'servicio_id')} — {texto(s['nombre_cliente'], 'nombre_cliente')}",
                        2)
        lineas = texto(s["lineas_asociadas"], "lineas_asociadas")
        if "," in lineas:
            parrafo(doc, f"{etiqueta('lineas_asociadas')}: {lineas}")
        pares = [(etiqueta(c), texto(s[c], c), color(c, texto(s[c], c))) for c in _METRICAS_FICHA]
        mitad = (len(pares) + 1) // 2
        filas, fondos = [], []
        for izq, der in zip(pares[:mitad], pares[mitad:] + [("", "", None)] * (2 * mitad - len(pares))):
            filas.append([izq[0], izq[1], der[0], der[1]])
            fondos.append([None, izq[2], None, der[2]])
        tabla(doc, ["Métrica", "Valor", "Métrica", "Valor"], filas, fondos, anchos=[2.2, 1.2, 2.2, 1.2],
              fraccion=0.8)

        parrafo(doc, "Reclamos del servicio (orden de atribución: fecha de inicio y número de reclamo)", _ESTILO_SUBTITULO)
        idx = self._reclamos_por_clave.get(s["clave_servicio"])
        if idx is None or len(idx) == 0:
            parrafo(doc, "Sin reclamos vinculados.")
        else:
            tabla(doc, [_ETIQUETAS_EXTRA.get(c) or etiqueta(c) for c in _RECLAMOS_FICHA],
                  [self._reclamos[i] for i in idx], [self._fondos_reclamos[i] for i in idx], _RECLAMOS_ANCHOS)

        parrafo(doc, "Eventos reales del servicio", _ESTILO_SUBTITULO)
        idx_ev = self._eventos_por_clave.get(s["clave_servicio"])
        if idx_ev is None or len(idx_ev) == 0:
            parrafo(doc, "Sin eventos reales.")
        else:
            tabla_df(doc, self._eventos.iloc[idx_ev], _EVENTOS_FICHA, _EVENTOS_ANCHOS)


def ficha_servicio(doc, res, fila_servicio) -> None:
    """Ficha completa de un servicio (para usos sueltos; en lote conviene `FichasServicio`)."""
    FichasServicio(res).escribir(doc, fila_servicio)
