# Nombre de archivo: xlsx_builder.py
# Ubicación de archivo: core/sla_consumo/xlsx_builder.py
# Descripción: XLSX completo del informe SLA consumido — una hoja por vista, con formatos y colores condicionales

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from core.sla_consumo.presentacion import (COLORES_ESTADO, COLORES_SEMAFORO, COLORES_SUFICIENTE, COLUMNAS,
                                           COLUMNAS_FRACCION, LEYENDAS, TOTALES, _faltante, _numero,
                                           a_fecha_ar, fmt_bool)

HOJAS = ("Resumen y leyendas", "Servicios", "Servicio x Reclamo", "Eventos", "Evento x Servicio",
         "Reclamos sin evento", "Inconsistencias", "No vinculados", "Codigo de cierre", "Detalle tipo solucion",
         "Carrier x Cliente")

NUMBER_FORMAT = {"horas": "[h]:mm", "pct": "0.00%", "fecha": "dd/mm/yyyy hh:mm", "entero": "0"}
_MAX_ANCHO = 50
_MAX_TEXTO = 32000  # límite de celda de Excel: 32767
_COLORES_CF = {"semaforo": COLORES_SEMAFORO, "indicador": COLORES_SEMAFORO, "estado": COLORES_ESTADO,
               "estado_servicio": COLORES_ESTADO, "suficiente_solo": COLORES_SUFICIENTE}
_CLARO = {"F1C40F", "BDC3C7"}   # fondos donde el texto blanco no se lee
_ENCABEZADO_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
_ENCABEZADO_FONT = Font(bold=True, color="FFFFFF")


def _texto(v: object) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", str(v))[:_MAX_TEXTO]


def _valor(v: object, formato: str, col: str = ""):
    """Convierte un valor del resultado al valor de celda; NaN/None/±inf → celda vacía; un texto queda texto."""
    if _faltante(v):
        return None
    if formato in ("horas", "pct", "entero"):
        n = _numero(v)
        if n is None:
            return _texto(v)
        if formato == "horas":
            return n / 24
        if formato == "pct":
            return n if col in COLUMNAS_FRACCION else n / 100
        return int(n)
    if formato == "fecha":
        return a_fecha_ar(v)
    if formato == "bool":
        return fmt_bool(v)
    return _texto(v)


def _columna_fecha(serie: pd.Series) -> list:
    return [a_fecha_ar(v) for v in serie]


def _ancho(etiqueta: str, valores: list, formato: str) -> float:
    if formato in ("horas", "pct", "entero", "bool"):
        largo_valores = 10
    elif formato == "fecha":
        largo_valores = 16
    else:
        largo_valores = max((len(str(v)) for v in valores[:500] if v is not None), default=0)
    return min(max(len(etiqueta) * 0.9, largo_valores) + 2, _MAX_ANCHO)


def _formato_condicional(ws, columna: str, letra: str, filas: int) -> None:
    rango = f"{letra}2:{letra}{max(filas + 1, 2)}"
    for texto, color in _COLORES_CF[columna].items():
        if color is None:
            continue
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        font = Font(color="000000" if color in _CLARO else "FFFFFF", bold=True)
        ws.conditional_formatting.add(
            rango, FormulaRule(formula=[f'${letra}2="{texto}"'], fill=fill, font=font))


def _hoja_datos(wb: Workbook, nombre: str, df: pd.DataFrame, renombres: dict[str, str] | None = None) -> None:
    ws = wb.create_sheet(nombre)
    renombres = renombres or {}
    claves = [renombres.get(c, c) for c in df.columns]
    desconocidas = [c for c in claves if c not in COLUMNAS]
    if desconocidas:
        raise KeyError(f"Columnas sin definición en presentacion.COLUMNAS: {desconocidas}")
    formatos = [COLUMNAS[c][1] for c in claves]
    ws.append([COLUMNAS[c][0] for c in claves])
    columnas = []
    for original, clave, formato in zip(df.columns, claves, formatos):
        serie = df[original]
        columnas.append(_columna_fecha(serie) if formato == "fecha"
                        else [_valor(v, formato, clave) for v in serie])
    for fila in zip(*columnas) if columnas else ():
        ws.append(list(fila))
    for j, (clave, formato) in enumerate(zip(claves, formatos), start=1):
        nf = NUMBER_FORMAT.get(formato)
        if nf:
            for (celda,) in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=j, max_col=j):
                celda.number_format = nf
        letra = get_column_letter(j)
        ws.column_dimensions[letra].width = _ancho(COLUMNAS[clave][0], columnas[j - 1], formato)
        if clave in _COLORES_CF:
            _formato_condicional(ws, clave, letra, len(df))
    for celda in ws[1]:
        celda.font, celda.fill = _ENCABEZADO_FONT, _ENCABEZADO_FILL
        celda.alignment = Alignment(wrap_text=True, vertical="center")
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "B2"


def _hoja_resumen(wb: Workbook, res) -> None:
    ws = wb.active
    ws.title = HOJAS[0]
    ws.append(["Concepto", "Valor", "% / Detalle"])
    negrita = Font(bold=True)
    secciones: list[int] = []
    formatos: dict[int, tuple[str | None, str | None]] = {}

    def fila(a, b=None, c=None, fb=None, fc=None, seccion=False):
        ws.append([a, b, c])
        n = ws.max_row
        formatos[n] = (fb, fc)
        if seccion:
            secciones.append(n)

    fila("Fecha de corte", dt.datetime.combine(res.fecha_corte, dt.time()), fb="dd/mm/yyyy")
    inicio = res.fecha_corte - dt.timedelta(days=365)
    fila("Ventana de análisis (365 días)", f"{inicio:%d/%m/%Y} a {res.fecha_corte:%d/%m/%Y}")
    fila("Totales (las horas, salvo las netas, cubren sólo reclamos vinculados)", seccion=True)
    for clave, (etiqueta, formato) in TOTALES.items():
        if clave in res.totales:
            fila(etiqueta, _valor(res.totales[clave], formato), fb=NUMBER_FORMAT.get(formato))
    fila("Composición de las horas computables por causa", "Horas", "% del total", seccion=True)
    for _, f in res.composicion.iterrows():
        fila(f["causa"], _valor(f["horas"], "horas"), _valor(f["pct"], "pct"), fb="[h]:mm", fc="0.00%")
    fila("Leyendas", "Explicación", seccion=True)
    for concepto, explicacion in LEYENDAS:
        fila(concepto, explicacion)
    for n in secciones:
        for celda in ws[n]:
            celda.font = negrita
    for n, (fb, fc) in formatos.items():
        if fb:
            ws.cell(n, 2).number_format = fb
        if fc:
            ws.cell(n, 3).number_format = fc
        ws.cell(n, 2).alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        ws.cell(n, 1).alignment = Alignment(wrap_text=True, vertical="top")
    for celda in ws[1]:
        celda.font, celda.fill = _ENCABEZADO_FONT, _ENCABEZADO_FILL
    ws.column_dimensions["A"].width = _MAX_ANCHO
    ws.column_dimensions["B"].width = _MAX_ANCHO
    ws.column_dimensions["C"].width = 18
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def construir_xlsx(res, destino: Path) -> Path:
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _hoja_resumen(wb, res)
    eventos_clave = {"eventos": "eventos_y_reclamos_sin_evento"}
    vistas = ((HOJAS[1], res.servicios, None), (HOJAS[2], res.servicio_reclamo, None), (HOJAS[3], res.eventos, None),
              (HOJAS[4], res.evento_servicio, None), (HOJAS[5], res.reclamos_sin_evento, None),
              (HOJAS[6], res.inconsistencias, None), (HOJAS[7], res.no_vinculados, None),
              (HOJAS[8], res.codigo_cierre, eventos_clave), (HOJAS[9], res.codigo_cierre_detalle, eventos_clave),
              (HOJAS[10], res.carrier_cliente, None))
    for nombre, df, renombres in vistas:
        _hoja_datos(wb, nombre, df, renombres)
    wb.save(destino)
    return destino
