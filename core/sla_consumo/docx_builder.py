# Nombre de archivo: docx_builder.py
# Ubicación de archivo: core/sla_consumo/docx_builder.py
# Descripción: DOCX ejecutivo del informe SLA consumido con gráficos de repercutores

from __future__ import annotations

import math
from pathlib import Path

from docx import Document
from docx.shared import Cm

from core.sla_consumo import charts


def _num(valor, formato: str = ",.2f") -> str:
    """Formatea un número; None/NaN/inf se muestran como '—'."""
    try:
        v = float(valor)
    except (TypeError, ValueError):
        return "—"
    return format(v, formato) if math.isfinite(v) else "—"


def _tabla(doc, df, columnas: list[tuple[str, str]], formato=lambda v: v) -> None:
    tabla = doc.add_table(rows=1, cols=len(columnas))
    tabla.style = "Light Grid Accent 1"
    for celda, (_, titulo) in zip(tabla.rows[0].cells, columnas):
        celda.text = titulo
    for _, fila in df.iterrows():
        celdas = tabla.add_row().cells
        for celda, (col, _) in zip(celdas, columnas):
            valor = fila[col]
            if isinstance(valor, float):
                celda.text = _num(valor)
            else:
                celda.text = "—" if valor is None or valor != valor else str(valor)


def construir_docx(res, destino: Path, top_servicios: int = 20) -> Path:
    destino.parent.mkdir(parents=True, exist_ok=True)
    graficos = destino.parent / f"{destino.stem}_graficos"
    doc = Document()
    doc.add_heading("Informe de SLA consumido", level=0)
    t = res.totales
    doc.add_paragraph(
        f"Fecha de corte: {res.fecha_corte.isoformat()} · ventana de 12 meses. "
        f"{t['reclamos']} reclamos en {t['eventos']} eventos (+{t['reclamos_aislados']} reclamos aislados), "
        f"{t['servicios']} servicios afectados, {t['servicios_excedidos']} superan su presupuesto anual de SLA. "
        f"Horas que consumen SLA: {t['horas_sla']:,.1f} de {t['horas']:,.1f} h netas.")

    doc.add_heading("SLA consumido por código de cierre", level=1)
    _tabla(doc, res.codigo_cierre, [("grupo", "Grupo"), ("eventos", "Eventos"), ("reclamos", "Reclamos"),
                                    ("horas", "Horas"), ("horas_sla", "Horas SLA"), ("pct_horas", "% horas")])
    doc.add_picture(str(charts.horas_por_grupo(res, graficos / "grupos.png")), width=Cm(15))

    doc.add_heading("Mayores repercutores — eventos", level=1)
    doc.add_picture(str(charts.pareto_eventos(res, graficos / "pareto_eventos.png")), width=Cm(16))
    _tabla(doc, res.eventos.head(15), [("evento_clave", "Evento"), ("reclamos", "Reclamos"),
                                       ("servicios_afectados", "Servicios"), ("servicios_excedidos", "Excedidos"),
                                       ("horas_sla", "Horas SLA"), ("grupo_predominante", "Grupo")])

    doc.add_heading("Mayores repercutores — servicios", level=1)
    doc.add_picture(str(charts.top_servicios(res, graficos / "top_servicios.png")), width=Cm(16))

    doc.add_heading("Fichas por servicio", level=1)
    for _, svc in res.servicios.dropna(subset=["pct_presupuesto"]).head(top_servicios).iterrows():
        doc.add_heading(f"{svc['numero_linea']} — {svc['nombre_cliente']}", level=2)
        doc.add_paragraph(
            f"SLA prometido {_num(svc['sla_prometido'], '.3g')} % · presupuesto {_num(svc['presupuesto_h'])} h · "
            f"consumido {_num(svc['horas_sla'])} h ({_num(svc['pct_presupuesto'], '.1f')} % del presupuesto, "
            f"{_num(svc['pp_disponibilidad'], '.3f')} pp) · "
            f"SLA calculado {_num(float(svc['sla_calculado']) * 100, '.3f')} %.")
        doc.add_picture(str(charts.servicio(res, svc["numero_linea"], graficos / f"svc_{svc['numero_linea']}.png")),
                        width=Cm(15))
    doc.save(destino)
    return destino
